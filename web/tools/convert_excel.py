import json
import sys
from pathlib import Path

import openpyxl

SHEET_NAME = "英文對話總表"


def read_sentences(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
    sentences = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        sentence_id = row[0]
        if sentence_id is None:
            continue
        sentences.append({
            "id": int(sentence_id),
            "category": row[1],
            "zh": row[2],
            "en": row[3],
        })
    return sentences


def write_js(sentences, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    json_text = json.dumps(sentences, ensure_ascii=False, indent=2)
    output_path.write_text(f"window.SENTENCES = {json_text};\n", encoding="utf-8")


def main():
    script_dir = Path(__file__).resolve().parent
    xlsx_path = script_dir.parent.parent / "app_data" / "1490句完整日常英文對話表.xlsx"
    output_path = script_dir.parent / "data" / "sentences.js"

    if not xlsx_path.exists():
        print(f"找不到 Excel 檔案: {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    sentences = read_sentences(xlsx_path)
    write_js(sentences, output_path)
    print(f"已轉換 {len(sentences)} 筆句子到 {output_path}")


if __name__ == "__main__":
    main()
